"""Microsoft Word (.docx) Parser for Text, Tables, and Embedded Images.

Extracts all paragraphs, structured table rows, and extracts embedded raster images
from the OpenXML package for visual inspection.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import List, Tuple
import cv2
import docx
import numpy as np

from app.core.cache_manager import CacheManager
from app.parsers.base_parser import BaseParser, DocumentMetadata, PageData, ProcessedDocument
from app.vision.image_enhancer import ImageEnhancer


class DocxParser(BaseParser):
    """Parser for .docx word documents with embedded image and table extraction."""

    def can_handle(self, file_path: Path) -> bool:
        """Checks if file is a Word document."""
        return file_path.suffix.lower() in [".docx", ".docm"]

    def parse(self, file_path: Path) -> ProcessedDocument:
        """Parses the Word document, extracting text, tables, and embedded images.

        Args:
            file_path: Path to the .docx file.

        Returns:
            ProcessedDocument containing extracted sections and images.
        """
        sha256 = CacheManager.compute_sha256(file_path)
        doc = docx.Document(file_path)

        # 1. Extract Paragraphs
        paragraph_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # 2. Extract Tables
        table_lines = []
        for table_idx, table in enumerate(doc.tables, start=1):
            table_lines.append(f"\n--- Tabla {table_idx} ---")
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_lines.append(" | ".join(row_cells))

        text_content = "\n".join(paragraph_texts)
        if table_lines:
            text_content += "\n" + "\n".join(table_lines)

        # 3. Extract Embedded Images
        extracted_images = self._extract_embedded_images(doc)

        pages: List[PageData] = []
        # Page 1: Text & Tables
        text_page = PageData(
            page_number=1,
            raw_text=text_content,
            cleaned_text=text_content,
            extra_metadata={"tables_count": len(doc.tables), "paragraphs_count": len(doc.paragraphs)},
        )
        pages.append(text_page)

        # Pages 2..N: Embedded images (if any)
        tree_detected = False
        has_signature = False
        for idx, (img_name, img_np) in enumerate(extracted_images, start=2):
            enhanced_img, blur_var, is_blurry = ImageEnhancer.enhance_document_page(img_np)
            img_page = PageData(
                page_number=idx,
                raw_text=f"[Imagen embebida: {img_name}]",
                cleaned_text=f"[Imagen embebida: {img_name}]",
                image_np=enhanced_img,
                blur_variance=blur_var,
                is_blurry=is_blurry,
                extra_metadata={"filename": img_name},
            )
            pages.append(img_page)

        # Build basic metadata (will be enriched by classifier / LLM if needed)
        metadata = DocumentMetadata(
            category="GENERAL",
            target_department="Administración General",
            confidence_score=0.7,
            tree_detected=tree_detected,
            has_signature_or_seal=has_signature,
            suggested_filename=f"DOC_{file_path.stem}.docx",
        )

        return ProcessedDocument(
            file_path=file_path,
            original_name=file_path.name,
            sha256_hash=sha256,
            file_type="DOCX",
            num_pages=len(pages),
            pages=pages,
            metadata=metadata,
            full_text=text_content,
            processing_time_sec=0.05,
        )

    def _extract_embedded_images(self, doc: docx.Document) -> List[Tuple[str, np.ndarray]]:
        """Extracts raster images embedded inside the DOCX zip package."""
        images: List[Tuple[str, np.ndarray]] = []
        try:
            for part in doc.part.related_parts.values():
                if "image" in part.content_type:
                    image_bytes = part.blob
                    filename = Path(part.partname).name
                    np_arr = np.frombuffer(image_bytes, np.uint8)
                    img = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)
                    if img is not None:
                        images.append((filename, img))
        except Exception:
            pass
        return images
